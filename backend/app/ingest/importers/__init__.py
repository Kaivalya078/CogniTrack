"""File formats. Each importer turns a file's bytes into proposed staging items and nothing else:
importers never touch the database (CLAUDE.md rule 7).

A lot item's payload (JSON, money as strings so it stays exact):
    strategy, account, symbol, isin, name, opened_on (ISO date or null), date_confidence,
    buys:    [{date, qty, price, charges, charges_estimated, source, source_ref}]
    exits:   [{label, date, qty, price, charges, charges_estimated, source, source_ref}]
    targets: [{label, price}]   planned, not yet hit
    stop, thesis, review_trigger, remarks,
    sheet (the file's own P&L and duration, to check), file_warnings
A cash item's payload: account, as_of, amount (both null until I fill them in).
"""

import io
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Protocol

import openpyxl
from openpyxl.utils.datetime import from_excel


@dataclass
class ProposedItem:
    row_ref: str
    kind: str  # "lot" or "cash"
    payload: dict[str, Any]


class Importer(Protocol):
    source: str

    def detect(self, data: bytes, filename: str) -> bool: ...
    def parse(self, data: bytes, filename: str) -> list[ProposedItem]: ...


def cash_item(account: str) -> ProposedItem:
    """Optional: left empty, it is simply not committed."""
    return ProposedItem("cash", "cash", {"account": account, "as_of": None, "amount": None})


def rupees(value: Any) -> Decimal:
    """An Excel number as exact rupees: 5794.400000000001 -> 5794.40."""
    return Decimal(str(value)).quantize(Decimal("0.01"), ROUND_HALF_UP)


def excel_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return from_excel(value).date()  # a serial number


def find_table(data: bytes, required: set[str]) -> tuple[str, int, list[str]] | None:
    """(sheet title, header row number, headers) of the first sheet with all `required`
    headers in one row, or None. Headers are stripped; columns are found by name, not position."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # not an xlsx at all
        return None
    for ws in wb.worksheets:
        for n, row in enumerate(ws.iter_rows(max_row=30, values_only=True), 1):
            headers = [str(v).strip() if v is not None else "" for v in row]
            if required <= set(headers):
                return ws.title, n, headers
    return None


def rows(data: bytes, required: set[str]) -> Iterator[tuple[str, dict[str, Any]]]:
    """(row ref like "Sheet1!A3", {header: value}) for each row below the header."""
    found = find_table(data, required)
    if found is None:
        raise ValueError(f"No sheet has the columns {sorted(required)}")
    title, header_row, headers = found
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for n, row in enumerate(wb[title].iter_rows(min_row=header_row + 1, values_only=True)):
        yield f"{title}!A{header_row + 1 + n}", dict(zip(headers, row, strict=False))
